from openpyxl import load_workbook

from order_parser.chat_parser import parse_chat_file
from order_parser.cli import run
from order_parser.detector import detect
from order_parser.report import buy_list, by_customer

SAMPLE = "examples/sample_chat.txt"
SELLER = "Magic Imports"


def sample_result():
    return detect(parse_chat_file(SAMPLE), [SELLER])


def test_sample_chat_buy_list_totals():
    rows = {r["code"]: r for r in buy_list(sample_result())}
    assert rows["A1"]["quantity"] == 3          # Carla's order was cancelled
    assert rows["A2"]["quantity"] == 3
    assert rows["A3"]["quantity"] == 4
    assert rows["A2"]["revenue_brl"] == 747.0   # 3 x 249,00


def test_same_customer_quantities_are_merged():
    rows = {r["code"]: r for r in buy_list(sample_result())}
    assert "Patrícia Alves (3)" in rows["B7"]["customers"]


def test_by_customer_totals():
    rows = {r["customer"]: r for r in by_customer(sample_result())}
    assert rows["Juliana Prado"]["units"] == 3
    assert rows["Juliana Prado"]["total_brl"] == round(2 * 189.90 + 249.00, 2)
    assert "Carla Mendes" in rows               # still has her Stitch order


def test_cli_writes_excel_with_three_sheets(tmp_path, capsys):
    out = tmp_path / "buy.xlsx"
    assert run([SAMPLE, "--seller", SELLER, "--out", str(out)]) == 0
    book = load_workbook(out)
    assert book.sheetnames == ["Buy list", "By customer", "All orders"]
    assert "BUY LIST" in capsys.readouterr().out


def test_cli_writes_csv_files(tmp_path):
    out = tmp_path / "buy.csv"
    run([SAMPLE, "--seller", SELLER, "--out", str(out)])
    assert (tmp_path / "buy_buy_list.csv").exists()
    assert (tmp_path / "buy_all_orders.csv").exists()


def test_since_filter_keeps_only_the_latest_round(capsys):
    run([SAMPLE, "--seller", SELLER, "--since", "17/09/2026 14:25"])
    out = capsys.readouterr().out
    assert "Kit Kat" in out and "Caneca" not in out


def test_wrong_seller_name_warns_and_lists_names(capsys):
    run([SAMPLE, "--seller", "Wrong Name"])
    err = capsys.readouterr().err
    assert "not found" in err and "Magic Imports" in err


def test_missing_file_returns_error():
    assert run(["nope.txt", "--seller", SELLER]) == 1
